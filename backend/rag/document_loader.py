"""文档加载器：解析 PDF / Word / 文本，清洗、分块。"""
import re
from pathlib import Path


def load_document_text(path) -> str:
    """根据扩展名解析文档，返回纯文本。"""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext in (".docx", ".doc"):
        return _load_docx(path)
    # txt / md / 其它：按 UTF-8 文本读取
    return path.read_text(encoding="utf-8", errors="ignore")


def _load_pdf(path: Path) -> str:
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise RuntimeError("未安装 PyPDF2，请执行 pip install PyPDF2 后重试") from exc
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _load_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def clean_text(text: str) -> str:
    """基础文本清洗。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_chunks(text: str, chunk_size: int = 500, overlap: int = 80) -> list:
    """文本分块：优先使用 LangChain 语义分块，否则按字符分块（中文友好）。"""
    text = clean_text(text)
    if not text:
        return []
    # 优先 LangChain RecursiveCharacterTextSplitter
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter  # type: ignore

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=overlap,
            separators=["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""],
        )
        chunks = [c.strip() for c in splitter.split_text(text) if c.strip()]
        if chunks:
            return chunks
    except Exception:
        pass
    # 兜底：按字符分块
    chunks = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        chunks.append(text[start:end].strip())
        if end >= n:
            break
        start = end - overlap
    return [c for c in chunks if c]


# 课标 / 公共教材类文档的特征关键词
PUBLIC_KEYWORDS = [
    "课程标准", "课程性质", "核心素养", "学段目标", "总目标", "内容要求",
    "学业质量", "课程理念", "教学建议", "课程目标", "教材", "考试说明",
]


def classify_document(text: str) -> str:
    """按内容自动分类文档：public（课标/教材类）或 private（教师个人材料）。"""
    head = (text or "")[:2000]
    score = sum(1 for k in PUBLIC_KEYWORDS if k in head)
    return "public" if score >= 2 else "private"
