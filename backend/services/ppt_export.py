"""把幻灯片 JSON 导出为 .pptx（需要 python-pptx，可选依赖）。"""
from io import BytesIO


def _slide_text(slide: dict) -> str:
    """从幻灯片 JSON 提取纯文本。"""
    t = slide.get("type")
    if t == "cover":
        return slide.get("title", "")
    if t == "objectives":
        return "\n".join(f"• {it}" for it in (slide.get("items") or []))
    if t == "content":
        cards = slide.get("cards") or []
        return (slide.get("heading", "") + "\n" + "\n".join(
            f"{c.get('title', '')}：{c.get('body', '')}" for c in cards
        )).strip()
    if t == "choice":
        return (slide.get("question", "") + "\n" + "\n".join(
            f"{chr(65 + i)}. {o}" for i, o in enumerate(slide.get("options") or [])
        ))
    if t == "judge":
        return "\n".join(f"• {s.get('text', '')}" for s in (slide.get("statements") or []))
    if t == "drag":
        return (slide.get("prompt", "") + "\n" + "\n".join(
            f"{i + 1}. {it}" for i, it in enumerate(slide.get("items") or [])
        ))
    if t == "fill":
        return slide.get("question", "")
    if t == "match":
        return (slide.get("question", "") + "\n" + "\n".join(
            f"{p.get('left', '')} —— {p.get('right', '')}" for p in (slide.get("pairs") or [])
        ))
    if t == "explore":
        return (slide.get("question", "") + "\n" + "\n".join(slide.get("poll") or []))
    if t == "summary":
        return "\n".join(f"• {p}" for p in (slide.get("points") or []))
    return ""


def deck_to_pptx(deck: dict) -> bytes:
    """把幻灯片内容渲染为 .pptx 字节。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    for s in deck.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # 仅标题版式
        title = slide.shapes.title
        if s.get("type") == "cover":
            title.text = deck.get("title", "交互式课件")
        else:
            title.text = s.get("type", "内容")
        body = _slide_text(s)
        if body:
            box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(9), Inches(5.5))
            tf = box.text_frame
            tf.text = body
            tf.word_wrap = True
            for p in tf.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(18)
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
