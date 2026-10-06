"""Markdown → Word(.docx) 导出。

把教案、习题、演讲稿等文本素材转成教师可直接用 Word 打开的 .docx 文件。
"""
import re
from io import BytesIO


def _add_runs(paragraph, text: str) -> None:
    """解析 **粗体** / *斜体* 内联标记。"""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _is_separator(cells) -> bool:
    return all(re.match(r"^:?-{2,}:?$", c) for c in cells)


def _add_table(doc, table_lines) -> None:
    rows = []
    for tl in table_lines:
        cells = [c.strip() for c in tl.strip().strip("|").split("|")]
        if _is_separator(cells):
            continue
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            table.rows[ri].cells[ci].text = row[ci] if ci < len(row) else ""


def markdown_to_docx(md_text: str) -> bytes:
    """把 Markdown 文本转换为 .docx 字节（需要安装 python-docx）。"""
    from docx import Document

    doc = Document()
    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        # 表格
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _add_table(doc, table_lines)
            continue
        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            i += 1
            continue
        # 无序列表
        if re.match(r"^[-*+]\s+", stripped):
            while i < len(lines) and re.match(r"^[-*+]\s+", lines[i].strip()):
                p = doc.add_paragraph(style="List Bullet")
                _add_runs(p, re.sub(r"^[-*+]\s+", "", lines[i].strip()))
                i += 1
            continue
        # 有序列表
        if re.match(r"^\d+[.)]\s+", stripped):
            while i < len(lines) and re.match(r"^\d+[.)]\s+", lines[i].strip()):
                p = doc.add_paragraph(style="List Number")
                _add_runs(p, re.sub(r"^\d+[.)]\s+", "", lines[i].strip()))
                i += 1
            continue
        # 分隔线 / 空行
        if re.match(r"^-{3,}$", stripped):
            doc.add_paragraph()
            i += 1
            continue
        # 普通段落
        _add_runs(doc.add_paragraph(), stripped)
        i += 1

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
